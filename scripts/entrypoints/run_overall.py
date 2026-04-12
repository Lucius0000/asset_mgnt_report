"""
"周报-资产大类表现-整体"表格计算。
"""

from __future__ import annotations

from pathlib import Path
import sys

PROJECT_ROOT = Path(__file__).resolve().parents[2]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import load_workbook

from src.asset_mgnt_report.metrics.sharpe import adjusted_sharpe_ratio, sharpe_ratio

CONFIG = {
    "input_path": Path("data") / "seeds" / "整体.xlsx",
    "output_path": Path("output") / "整体_processed.xlsx",
    "log_path": Path("output") / "raw_data" / "整体_calculation_steps.txt",
}


def read_percent_as_fraction(cell):
    v = cell.value
    if v is None:
        return None
    if isinstance(v, str):
        s = v.strip().replace("％", "%")
        if s.endswith("%"):
            s = s[:-1]
        return float(s) / 100.0
    fmt = (cell.number_format or "").lower()
    return float(v) if "%" in fmt else float(v) / 100.0


def plain_number_from_percent_cell(cell):
    v = cell.value
    if v is None:
        return None
    if isinstance(v, str):
        s = v.strip().replace("％", "%")
        if s.endswith("%"):
            s = s[:-1]
        return round(float(s), 2)
    fmt = (cell.number_format or "").lower()
    return round(float(v) * 100.0, 2) if "%" in fmt else round(float(v), 2)


def find_yoy(sheet, pair_name):
    header = [sheet.cell(1, c).value for c in range(1, sheet.max_column + 1)]
    yoy_cols = [i for i, v in enumerate(header, start=1) if v and "yoy" in str(v).lower()]
    for r in range(2, sheet.max_row + 1):
        row_has_pair = any(sheet.cell(r, c).value == pair_name for c in range(1, sheet.max_column + 1))
        if not row_has_pair:
            continue
        for yc in yoy_cols:
            val = sheet.cell(r, yc).value
            if val is not None:
                return float(val) / 100.0
    return None


def fmt_frac(x, d=6):
    return "None" if x is None else f"{x:.{d}f}"


def fmt_pctnum(x, d=2):
    if x is None:
        return "None"
    return f"{x * 100:.{d}f}"


def main() -> None:
    input_path = CONFIG["input_path"]
    output_path = CONFIG["output_path"]
    log_path = CONFIG["log_path"]
    wb = load_workbook(input_path)
    ws = wb["Sheet1"]
    ws2 = wb["Sheet2"]

    rf_by_region = {}
    current_region = None
    for r in range(2, ws.max_row + 1):
        region_cell_val = ws.cell(r, 1).value
        if region_cell_val not in (None, ""):
            current_region = str(region_cell_val).strip()
        asset = ws.cell(r, 2).value
        ret_frac = read_percent_as_fraction(ws.cell(r, 3))
        if asset == "债券固收" and current_region and ret_frac is not None:
            rf_by_region[current_region] = ret_frac

    us_rf = rf_by_region.get("美国", 0.0)
    cn_rf = rf_by_region.get("中国", None)
    hk_rf = 0.0062
    usd_cnh_yoy = find_yoy(ws2, "USD_CNH")
    usd_hkd_yoy = find_yoy(ws2, "USD_HKD")

    log_lines = []

    def log(message: str) -> None:
        print(message)
        log_lines.append(message)

    log("=== Calculation Steps ===")
    current_region = None
    for r in range(2, ws.max_row + 1):
        region_cell_val = ws.cell(r, 1).value
        if region_cell_val not in (None, ""):
            current_region = str(region_cell_val).strip()
        region = current_region
        asset = ws.cell(r, 2).value
        ret = read_percent_as_fraction(ws.cell(r, 3))
        vol = read_percent_as_fraction(ws.cell(r, 5))
        if ret is None or vol is None:
            continue
        if asset == "债券固收":
            ws.cell(r, 7).value = 0.00
            ws.cell(r, 8).value = 0.00
            continue

        rf_local = {"美国": us_rf, "中国": cn_rf, "香港": hk_rf}.get(region, us_rf)
        if rf_local is None:
            rf_local = us_rf
        sharpe = sharpe_ratio(ret, vol, rf_local)

        currency_effect = 0.0
        if region == "中国" and usd_cnh_yoy is not None:
            currency_effect = -usd_cnh_yoy
        elif region == "香港" and usd_hkd_yoy is not None:
            currency_effect = -usd_hkd_yoy
        adj_sharpe = adjusted_sharpe_ratio(ret, vol, us_rf, currency_effect)
        if asset in ("商品与贵金属（黄金）", "数字货币（BTC）"):
            adj_sharpe = sharpe

        log(
            f"[Row {r}] region={region} asset={asset} return={fmt_pctnum(ret)} "
            f"vol={fmt_pctnum(vol)} sharpe={fmt_frac(sharpe)} adjusted={fmt_frac(adj_sharpe)}"
        )
        ws.cell(r, 7).value = round(sharpe, 2)
        ws.cell(r, 7).number_format = "0.00"
        ws.cell(r, 8).value = round(adj_sharpe, 2)
        ws.cell(r, 8).number_format = "0.00"

    for r in range(2, ws.max_row + 1):
        for c in (3, 4, 5):
            cell = ws.cell(r, c)
            plain = plain_number_from_percent_cell(cell)
            if plain is not None:
                cell.value = plain
                cell.number_format = "0.00"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    log_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    log_path.write_text("\n".join(log_lines), encoding="utf-8")


if __name__ == "__main__":
    main()
