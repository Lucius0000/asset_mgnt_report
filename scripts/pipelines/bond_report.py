"""
债券报表：输出 bonds.xlsx（转置），并拆分“总市值 ($)”为“国债总市值 ($)”与“债券市场总市值 ($)”。

新增能力：
- 支持通过函数参数或 CLI 指定单次日期。
- 支持通过起止日期补录区间内的全部周五周报。
- 直接单独运行脚本时，自动补齐项目根目录到 sys.path，避免 `ModuleNotFoundError: src`。
"""

from __future__ import annotations

import argparse
import logging
import os
import sys
from datetime import datetime, timedelta
from io import StringIO
from pathlib import Path
from typing import Optional

import akshare as ak
import numpy as np
import pandas as pd
import requests
from fredapi import Fred
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src.asset_mgnt_report.config.defaults import build_app_config
from src.asset_mgnt_report.config.inputs import parse_bool, resolve_config_value
from src.asset_mgnt_report.io.overall_seed_snapshot import upsert_sheet1_asset_metrics


RAW_DIR = ROOT / "output" / "raw_data"
RAW_DIR.mkdir(parents=True, exist_ok=True)
APP_CONFIG = build_app_config(project_root=ROOT)
logger = logging.getLogger(__name__)

# 顶部配置区（适合 Spyder 直接运行）
# - debug: 布尔值，True / False
# - bond_date: 单次运行日期，格式 YYYY-MM-DD 或 YYYYMMDD，例如 2026-03-27
# - start_date / end_date: 区间补录日期，格式 YYYY-MM-DD 或 YYYYMMDD；脚本会自动筛出区间内周五
# - write_history / update_snapshot: 布尔值，True / False
#   write_history=True 时会在 output 根目录额外生成 bonds_YYYYMMDD.xlsx，不会创建 history 子目录
CONFIG = {
    "debug": False,
    "bond_date": '2026-04-10',
    "start_date": None,
    "end_date": None,
    "write_history": False,
    "update_snapshot": True,
}


def _format_billion(value_billion: float, currency: str) -> str:
    return f"{value_billion:,.0f} B {currency}"


def _parse_percent_text(value):
    if value in (None, "", "-"):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("％", "%")
    if text.endswith("%"):
        text = text[:-1]
    return float(text)


def _most_recent_friday(ref_dt: Optional[datetime] = None) -> datetime:
    ref_dt = ref_dt or datetime.today()
    days_back = (ref_dt.weekday() - 4) % 7
    return ref_dt - timedelta(days=days_back)


def _to_yyyymmdd(d: datetime) -> str:
    return d.strftime("%Y%m%d")


def _parse_cli_date(value: str | None) -> str | None:
    if value in (None, ""):
        return None
    text = str(value).strip()
    for fmt in ("%Y%m%d", "%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).strftime("%Y%m%d")
        except ValueError:
            continue
    raise ValueError(f"无效日期格式：{value}，请使用 YYYYMMDD 或 YYYY-MM-DD")


def _iter_fridays(start_date: str, end_date: str) -> list[str]:
    start_dt = datetime.strptime(start_date, "%Y%m%d")
    end_dt = datetime.strptime(end_date, "%Y%m%d")
    if start_dt > end_dt:
        raise ValueError("开始日期不能晚于结束日期")
    cursor = start_dt
    dates: list[str] = []
    while cursor <= end_dt:
        if cursor.weekday() == 4:
            dates.append(cursor.strftime("%Y%m%d"))
        cursor += timedelta(days=1)
    return dates


def _resolve_sse_date(explicit_date: str | None = None) -> str:
    """解析上交所债券日期。

    优先顺序：
    1. 显式传入参数
    2. Spyder/IPython 控制台变量
    3. 环境变量 `AK_BOND_DATE`
    4. 交互输入
    5. 昨天
    """
    if explicit_date:
        return explicit_date

    try:
        from IPython import get_ipython  # type: ignore

        ip = get_ipython()
        if ip is not None:
            ns = getattr(ip, "user_ns", {}) or {}
            for key in ("ak_bond_date", "bond_date", "sse_date"):
                value = ns.get(key)
                if isinstance(value, str) and value.isdigit() and len(value) == 8:
                    return value
    except Exception:
        pass

    value = os.getenv("AK_BOND_DATE")
    if isinstance(value, str) and value.isdigit() and len(value) == 8:
        return value

    try:
        user_in = input("请输入数据收盘日期，需要周中交易日，回车则选择昨天，yyyymmdd：").strip()
        if user_in and user_in.isdigit() and len(user_in) == 8:
            return user_in
    except Exception:
        pass

    return (datetime.today() - timedelta(days=1)).strftime("%Y%m%d")


def _get_china_market_caps(date_yyyymmdd: str) -> tuple[str, str, pd.DataFrame]:
    df = ak.bond_cash_summary_sse(date=date_yyyymmdd)
    df.to_excel(RAW_DIR / f"国债总市值_{date_yyyymmdd}.xlsx", index=False)

    working = df.copy()
    key_col = "债券现货" if "债券现货" in working.columns else working.columns[0]
    value_col = "托管市值" if "托管市值" in working.columns else "托管面值"

    treasury_row = working[working[key_col].astype(str).str.contains("国债", na=False)]
    total_row = working[working[key_col].astype(str).str.contains("合计", na=False)]

    treasury_val_b = float(treasury_row[value_col].iloc[0]) / 10 if not treasury_row.empty else np.nan
    total_val_b = float(total_row[value_col].iloc[0]) / 10 if not total_row.empty else np.nan
    return _format_billion(treasury_val_b, "CNY"), _format_billion(total_val_b, "CNY"), working


def _get_china_volumes_30d(ref_date_yyyymmdd: str | None = None) -> tuple[Optional[str], Optional[str]]:
    ref_date = datetime.strptime(ref_date_yyyymmdd, "%Y%m%d") if ref_date_yyyymmdd else datetime.today()
    start_date = ref_date - timedelta(days=30)
    date_range = pd.date_range(start=start_date, end=ref_date)

    all_frames = []
    for current in date_range:
        date_str = current.strftime("%Y%m%d")
        try:
            df = ak.bond_deal_summary_sse(date=date_str)
            if df is not None and not df.empty:
                df["数据日期"] = current.strftime("%Y-%m-%d")
                all_frames.append(df)
        except Exception:
            continue

    if not all_frames:
        return None, None

    full_df = pd.concat(all_frames, ignore_index=True)
    try:
        full_df.to_excel(RAW_DIR / f"all_bond_deal_data_{ref_date.strftime('%Y%m%d')}.xlsx", index=False)
    except Exception:
        pass

    latest_day = full_df["数据日期"].max()
    latest_record = full_df[(full_df["数据日期"] == latest_day) & (full_df["债券类型"] == "记账式国债")]
    day_amt_b = None
    if not latest_record.empty:
        day_amt_b = _format_billion(float(latest_record["当日成交金额"].values[0]) / 100000, "CNY")

    group_sum_df = full_df.groupby("债券类型")["当日成交金额"].sum().reset_index()
    month_amt_b = None
    if not group_sum_df.empty and (group_sum_df["债券类型"] == "记账式国债").any():
        total_amt = float(group_sum_df.loc[group_sum_df["债券类型"] == "记账式国债", "当日成交金额"].values[0])
        month_amt_b = _format_billion(total_amt / 100000, "CNY")

    return day_amt_b, month_amt_b


def _get_cn_us_yield_metrics(ref_date_yyyymmdd: str | None = None):
    def compute_cn(df: pd.DataFrame, col_name: str):
        x = df[["日期", col_name]].dropna().copy()
        x["日期"] = pd.to_datetime(x["日期"])
        x.set_index("日期", inplace=True)
        x.sort_index(inplace=True)
        x["r_daily"] = (1 + x[col_name] / 100) ** (1 / 252) - 1
        end_date = x.index.max()
        recent_30d = x.loc[end_date - timedelta(days=30): end_date]
        r30 = (1 + recent_30d["r_daily"]).prod() ** (252 / max(1, len(recent_30d))) - 1
        recent_1y = x.loc[end_date - timedelta(days=365): end_date]
        r1y = (1 + recent_1y["r_daily"]).prod() - 1
        vol_a = x.iloc[-252:]["r_daily"].std() * np.sqrt(252)
        sharpe = (r1y - 0.017) / vol_a if vol_a != 0 else np.nan
        return round(r30 * 100, 2), round(r1y * 100, 2), round(vol_a * 100, 2), round(sharpe, 2)

    def compute_us(df: pd.DataFrame, col: str):
        x = df[[col]].dropna().copy()
        x.index = pd.to_datetime(x.index)
        x.sort_index(inplace=True)
        x["r_daily"] = (1 + x[col]) ** (1 / 252) - 1
        end_date = x.index.max()
        recent_30d = x.loc[end_date - timedelta(days=30): end_date]
        r30 = (1 + recent_30d["r_daily"]).prod() ** (252 / max(1, len(recent_30d))) - 1
        recent_1y = x.loc[end_date - timedelta(days=365): end_date]
        r1y = (1 + recent_1y["r_daily"]).prod() - 1
        vol_a = recent_1y["r_daily"].std() * np.sqrt(252)
        sharpe = (r1y - 0.045) / vol_a if vol_a != 0 else np.nan
        return round(r30 * 100, 2), round(r1y * 100, 2), round(vol_a * 100, 2), round(sharpe, 2)

    cn_us = ak.bond_zh_us_rate(start_date="20240101")

    if ref_date_yyyymmdd and "日期" in cn_us.columns:
        ref_ts = pd.Timestamp(datetime.strptime(ref_date_yyyymmdd, "%Y%m%d").date())
        dated = cn_us.copy()
        dated["日期"] = pd.to_datetime(dated["日期"])
        filtered = dated[dated["日期"] <= ref_ts]
        if not filtered.empty:
            cn_us = filtered

    def _pick_yield_col(df: pd.DataFrame, country: str, tenor: str) -> str:
        tenor_candidates = [
            f"{country}国债收益率:{tenor}",
            f"{country}国债收益率{tenor}",
            f"{country}国债收益率 {tenor}",
        ]
        for column in df.columns:
            for tenor_candidate in tenor_candidates:
                if tenor_candidate in str(column):
                    return column
        for column in df.columns:
            text = str(column)
            if country in text and "国债" in text and (tenor.replace("年", "") in text):
                return column
        raise KeyError(f"未找到{country}国债收益率列：{tenor}")

    cn2 = compute_cn(cn_us, _pick_yield_col(cn_us, "中国", "2年"))
    cn10 = compute_cn(cn_us, _pick_yield_col(cn_us, "中国", "10年"))

    def _normalize_us_yield_frame(df: pd.DataFrame) -> pd.DataFrame:
        normalized = df.copy()
        normalized.index = pd.to_datetime(normalized.index)
        normalized = normalized[["2Y", "10Y"]].apply(pd.to_numeric, errors="coerce")
        if normalized[["2Y", "10Y"]].max(numeric_only=True).max() > 1:
            normalized = normalized / 100.0
        normalized.sort_index(inplace=True)
        return normalized

    def _load_cached_us_yield_frame(suffix: str) -> pd.DataFrame | None:
        exact = RAW_DIR / f"美债收益率_{suffix}.xlsx"
        candidates = [exact] if exact.exists() else []
        candidates.extend(sorted(RAW_DIR.glob("美债收益率_*.xlsx"), reverse=True))
        legacy = RAW_DIR / "美债收益率.xlsx"
        if legacy.exists():
            candidates.append(legacy)

        seen: set[Path] = set()
        for path in candidates:
            if path in seen:
                continue
            seen.add(path)
            try:
                cached = pd.read_excel(path, index_col=0)
                if {"2Y", "10Y"}.issubset(cached.columns):
                    logger.warning("美债收益率实时数据拉取失败，改用缓存文件：%s", path.name)
                    return _normalize_us_yield_frame(cached)
            except Exception:
                continue
        return None

    def _build_us_yield_frame_from_akshare(df: pd.DataFrame) -> pd.DataFrame:
        working = df.copy()
        working["日期"] = pd.to_datetime(working["日期"])
        working.set_index("日期", inplace=True)
        return _normalize_us_yield_frame(
            pd.DataFrame(
                {
                    "2Y": working[_pick_yield_col(df, "美国", "2年")],
                    "10Y": working[_pick_yield_col(df, "美国", "10年")],
                },
                index=working.index,
            )
        )

    def _fetch_us_yield_frame_from_fred(end_date_text: str) -> pd.DataFrame:
        codes = {"2Y": "DGS2", "10Y": "DGS10"}
        data = {}
        for key, code in codes.items():
            url = (
                "https://fred.stlouisfed.org/graph/fredgraph.csv"
                f"?id={code}&cosd=2024-01-01&coed={end_date_text}"
            )
            response = requests.get(url, timeout=15)
            response.raise_for_status()
            fred_df = pd.read_csv(StringIO(response.text))
            fred_df["observation_date"] = pd.to_datetime(fred_df["observation_date"])
            series = pd.to_numeric(fred_df[code].replace(".", np.nan), errors="coerce")
            data[key] = pd.Series(series.to_numpy(), index=fred_df["observation_date"]) / 100.0
        return _normalize_us_yield_frame(pd.DataFrame(data))

    end_date = datetime.strptime(ref_date_yyyymmdd, "%Y%m%d").strftime("%Y-%m-%d") if ref_date_yyyymmdd else datetime.today().strftime("%Y-%m-%d")
    suffix = ref_date_yyyymmdd or datetime.today().strftime("%Y%m%d")
    try:
        us_yield_df = _build_us_yield_frame_from_akshare(cn_us)
    except Exception as akshare_exc:
        try:
            logger.warning("AkShare 美债收益率列不可用，尝试 FRED 官方 CSV：%s", akshare_exc)
            us_yield_df = _fetch_us_yield_frame_from_fred(end_date)
        except Exception as fred_exc:
            cached = _load_cached_us_yield_frame(suffix)
            if cached is None:
                raise RuntimeError("美债收益率数据拉取失败，且未找到可用缓存文件。") from fred_exc
            us_yield_df = cached
    us2 = compute_us(us_yield_df, "2Y")
    us10 = compute_us(us_yield_df, "10Y")
    try:
        us_yield_df.index.name = "Date"
        us_yield_df.to_excel(RAW_DIR / f"美债收益率_{suffix}.xlsx")
    except Exception:
        pass
    return cn2, cn10, us2, us10


def _get_us_market_caps() -> tuple[Optional[str], Optional[str]]:
    treasury_b = None
    csv_files = sorted((ROOT / "data").glob("MSPD_SumSecty*.csv"))
    if csv_files:
        frames = []
        for fp in csv_files:
            try:
                df = pd.read_csv(fp)
                required = {"Record Date", "Security Type Description", "Total Public Debt Outstanding (in Millions)"}
                if required.issubset(df.columns):
                    frames.append(df)
            except Exception:
                continue
        if frames:
            merged = pd.concat(frames, ignore_index=True)
            merged["Record Date"] = pd.to_datetime(merged["Record Date"], errors="coerce")
            merged = merged[merged["Security Type Description"].astype(str) == "Total Marketable"].dropna(subset=["Record Date"])
            if not merged.empty:
                latest_row = merged.sort_values("Record Date").iloc[-1]
                millions = float(latest_row["Total Public Debt Outstanding (in Millions)"])
                treasury_b = _format_billion(millions / 1000.0, "USD")

    market_b = None
    try:
        fred = Fred(api_key=os.environ["FRED_API_KEY"])
        series = fred.get_series("GFDEBTN")
        market_b = _format_billion(float(series.iloc[-1]) / 1000.0, "USD")
    except Exception:
        pass

    return treasury_b, market_b


def _build_report_frame(
    *,
    sse_date: str,
) -> tuple[pd.DataFrame, dict[str, float | str], dict[str, float | str]]:
    cn_treasury_b, cn_market_b, _ = _get_china_market_caps(sse_date)
    cn_day_vol_b, cn_month_vol_b = _get_china_volumes_30d(sse_date)
    (cn2_r30, cn2_r1y, cn2_vol, _), (cn10_r30, cn10_r1y, cn10_vol, _), (us2_r30, us2_r1y, us2_vol, _), (us10_r30, us10_r1y, us10_vol, _) = _get_cn_us_yield_metrics(sse_date)
    us_treasury_b, us_market_b = _get_us_market_caps()

    header = [
        "指标类别", "指标", "种类", "月收益率年化 (%)", "年收益率 (%)", "年化波动率 (%)",
        "国债总市值 ($)", "债券市场总市值 ($)", "当日交易量", "月交易量",
    ]
    rows = [
        ["中国", "记账式国债", "2年期", f"{cn2_r30:.2f}%", f"{cn2_r1y:.2f}%", f"{cn2_vol:.2f}%", cn_treasury_b, cn_market_b, cn_day_vol_b or "-", cn_month_vol_b or "-"],
        [None, None, "10年期", f"{cn10_r30:.2f}%", f"{cn10_r1y:.2f}%", f"{cn10_vol:.2f}%", "-", "-", "-", "-"],
        [None, "储蓄式国债", "3年期", "-", "-", "-", "-", "-", "-", "-"],
        [None, None, "5年期", "-", "-", "-", "-", "-", "-", "-"],
        ["美国", "记账式国债", "2年期", f"{us2_r30:.2f}%", f"{us2_r1y:.2f}%", f"{us2_vol:.2f}%", us_treasury_b or "-", us_market_b or "-", "-", "-"],
        [None, None, "10年期", f"{us10_r30:.2f}%", f"{us10_r1y:.2f}%", f"{us10_vol:.2f}%", "-", "-", "-", "-"],
        [None, "储蓄式国债", "EE bonds", "-", "-", "-", "-", "-", "-", "-"],
        [None, None, "I bonds", "-", "-", "-", "-", "-", "-", "-"],
    ]

    report_df = pd.DataFrame(rows, columns=header).fillna("-")

    snapshot_payload_cn: dict[str, float | str] = {
        "月收益率年化 (%)": round(cn2_r30 / 100.0, 6),
        "年收益率 (%)": round(cn2_r1y / 100.0, 6),
        "月波动率年化（%）": round(cn2_vol / 100.0, 6),
        "总市值 ($)": cn_market_b,
    }
    snapshot_payload_us: dict[str, float | str] = {
        "月收益率年化 (%)": round(us2_r30 / 100.0, 6),
        "年收益率 (%)": round(us2_r1y / 100.0, 6),
        "月波动率年化（%）": round(us2_vol / 100.0, 6),
        "总市值 ($)": us_market_b or "-",
    }
    return report_df, snapshot_payload_cn, snapshot_payload_us


def _write_bond_workbook(df_t: pd.DataFrame, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df_t.to_excel(writer, sheet_name="bonds", header=False, index=True, startrow=0)
        ws = writer.sheets["bonds"]

        max_row = ws.max_row
        max_col = ws.max_column
        for row_idx in range(1, max_row + 1):
            for col_idx in range(1, max_col + 1):
                ws.cell(row=row_idx, column=col_idx).alignment = Alignment(horizontal="center", vertical="center")

        cn_start, cn_end = 2, 5
        us_start, us_end = 6, 9
        ws.merge_cells(start_row=1, start_column=cn_start, end_row=1, end_column=cn_end)
        ws.merge_cells(start_row=1, start_column=us_start, end_row=1, end_column=us_end)
        ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=3)
        ws.merge_cells(start_row=2, start_column=4, end_row=2, end_column=5)
        ws.merge_cells(start_row=2, start_column=6, end_row=2, end_column=7)
        ws.merge_cells(start_row=2, start_column=8, end_row=2, end_column=9)
        for row_idx in (7, 8):
            ws.merge_cells(start_row=row_idx, start_column=cn_start, end_row=row_idx, end_column=cn_end)
            ws.merge_cells(start_row=row_idx, start_column=us_start, end_row=row_idx, end_column=us_end)
        for row_idx in (4, 5, 6):
            for col_idx in range(2, max_col + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                parsed = _parse_percent_text(cell.value)
                if parsed is None:
                    continue
                cell.value = parsed / 100.0
                cell.number_format = "0.00%"
        for column_cells in ws.columns:
            column_letter = get_column_letter(column_cells[0].column)
            max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
            ws.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 28)


def _run_for_single_date(
    *,
    sse_date: str,
    update_snapshot: bool = True,
    write_latest_output: bool = True,
    write_history_output: bool = False,
) -> list[Path]:
    os.environ["AK_BOND_DATE"] = sse_date

    report_df, snapshot_payload_cn, snapshot_payload_us = _build_report_frame(sse_date=sse_date)
    if update_snapshot:
        upsert_sheet1_asset_metrics(
            region="中国",
            asset_class="债券固收",
            fields=snapshot_payload_cn,
            config=APP_CONFIG,
        )
        upsert_sheet1_asset_metrics(
            region="美国",
            asset_class="债券固收",
            fields=snapshot_payload_us,
            config=APP_CONFIG,
        )

    df_t = report_df.T
    outputs: list[Path] = []
    if write_latest_output:
        latest_output = ROOT / "output" / "bonds.xlsx"
        _write_bond_workbook(df_t, latest_output)
        outputs.append(latest_output)
    if write_history_output:
        history_output = ROOT / "output" / f"bonds_{sse_date}.xlsx"
        _write_bond_workbook(df_t, history_output)
        outputs.append(history_output)
    return outputs


def main(
    debug: bool | None = None,
    bond_date: str | None = None,
    start_date: str | None = None,
    end_date: str | None = None,
    write_history: bool | None = None,
    update_snapshot: bool | None = None,
):
    """生成债券周报。

    `bond_date` 用于单次运行。
    `start_date` + `end_date` 用于补录区间内的所有周五。
    """

    resolved_debug = resolve_config_value(
        explicit=debug,
        env_key="AMR_BOND_DEBUG",
        default=CONFIG["debug"],
        caster=parse_bool,
    )
    resolved_bond_date = resolve_config_value(
        explicit=bond_date,
        env_key="AMR_BOND_DATE",
        default=CONFIG["bond_date"],
        caster=lambda raw: raw,
    )
    resolved_start_date = resolve_config_value(
        explicit=start_date,
        env_key="AMR_BOND_START_DATE",
        default=CONFIG["start_date"],
        caster=lambda raw: raw,
    )
    resolved_end_date = resolve_config_value(
        explicit=end_date,
        env_key="AMR_BOND_END_DATE",
        default=CONFIG["end_date"],
        caster=lambda raw: raw,
    )
    resolved_write_history = resolve_config_value(
        explicit=write_history,
        env_key="AMR_BOND_WRITE_HISTORY",
        default=CONFIG["write_history"],
        caster=parse_bool,
    )
    resolved_update_snapshot = resolve_config_value(
        explicit=update_snapshot,
        env_key="AMR_BOND_UPDATE_SNAPSHOT",
        default=CONFIG["update_snapshot"],
        caster=parse_bool,
    )

    explicit_bond_date = _parse_cli_date(resolved_bond_date)
    explicit_start_date = _parse_cli_date(resolved_start_date)
    explicit_end_date = _parse_cli_date(resolved_end_date)

    if explicit_bond_date and (explicit_start_date or explicit_end_date):
        raise ValueError("bond_date 与 start_date/end_date 不能同时使用")
    if explicit_start_date and not explicit_end_date:
        explicit_end_date = explicit_start_date
    if explicit_end_date and not explicit_start_date:
        explicit_start_date = explicit_end_date

    if resolved_debug is False and explicit_bond_date is None and explicit_start_date is None:
        explicit_bond_date = _to_yyyymmdd(_most_recent_friday())
    elif resolved_debug is True and explicit_bond_date is None and explicit_start_date is None:
        os.environ.pop("AK_BOND_DATE", None)

    if explicit_start_date:
        friday_dates = _iter_fridays(explicit_start_date, explicit_end_date or explicit_start_date)
        if not friday_dates:
            raise ValueError("给定日期区间内没有周五，无法补录周报数据")
        created_files: list[Path] = []
        for idx, sse_date in enumerate(friday_dates):
            created_files.extend(
                _run_for_single_date(
                    sse_date=sse_date,
                    update_snapshot=bool(resolved_update_snapshot) and idx == len(friday_dates) - 1,
                    write_latest_output=idx == len(friday_dates) - 1,
                    write_history_output=True,
                )
            )
        print("已生成以下债券周报文件：")
        for path in created_files:
            print(f"- {path.name}")
        return

    sse_date = _resolve_sse_date(explicit_bond_date)
    created_files = _run_for_single_date(
        sse_date=sse_date,
        update_snapshot=bool(resolved_update_snapshot),
        write_latest_output=True,
        write_history_output=bool(resolved_write_history),
    )
    print("已生成以下债券周报文件：")
    for path in created_files:
        print(f"- {path.name}")


def _build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="生成债券周报，支持单次指定日期和区间补录。")
    parser.add_argument("--debug", action="store_true", default=None, help="启用交互模式；未传日期时提示输入。")
    parser.add_argument("--bond-date", help="指定单次运行日期，支持 YYYYMMDD 或 YYYY-MM-DD。")
    parser.add_argument("--start-date", help="指定补录开始日期，支持 YYYYMMDD 或 YYYY-MM-DD。")
    parser.add_argument("--end-date", help="指定补录结束日期，支持 YYYYMMDD 或 YYYY-MM-DD。")
    parser.add_argument("--write-history", action="store_true", default=None, help="单次运行时在 output 根目录额外写出 bonds_YYYYMMDD.xlsx。")
    parser.add_argument("--skip-snapshot-update", action="store_true", default=None, help="只生成 Excel，不更新 overall_seed_snapshot。")
    return parser


if __name__ == "__main__":
    args = _build_arg_parser().parse_args()
    main(
        debug=args.debug,
        bond_date=args.bond_date,
        start_date=args.start_date,
        end_date=args.end_date,
        write_history=args.write_history,
        update_snapshot=None if args.skip_snapshot_update is None else (not args.skip_snapshot_update),
    )
