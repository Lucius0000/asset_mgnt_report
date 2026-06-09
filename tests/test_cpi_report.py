from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from scripts.pipelines import bond_report, cpi_report, crypto_report


def _sample_standard_frame(source: str = "source") -> pd.DataFrame:
    dates = pd.date_range("2014-01-01", periods=150, freq="MS")
    index_values = pd.Series(range(100, 250), dtype=float)
    frame = pd.DataFrame({"日期": dates, "指数": index_values})
    frame["MoM"] = frame["指数"].pct_change() * 100
    frame["YoY"] = frame["指数"].pct_change(12) * 100
    frame["数据源"] = source
    return frame


def test_get_cpi_data_falls_back_to_akshare_for_china(monkeypatch) -> None:
    us_frame = _sample_standard_frame("fred")
    hk_frame = _sample_standard_frame("hk_censtatd_api")
    cn_frame = _sample_standard_frame("china_akshare_macro_china_cpi")

    monkeypatch.setattr(cpi_report, "_fetch_us_series_with_fred", lambda *args, **kwargs: us_frame.copy())
    monkeypatch.setattr(cpi_report, "_fetch_china_cpi_from_nbs_official", lambda: (_ for _ in ()).throw(RuntimeError("blocked")))
    monkeypatch.setattr(cpi_report, "_fetch_china_cpi_from_akshare", lambda: cn_frame.copy())
    monkeypatch.setattr(cpi_report, "_fetch_hk_cpi_from_api", lambda *args, **kwargs: hk_frame.copy())

    data = cpi_report.get_cpi_data(time_range=24)

    assert data["CN_CPI"]["数据源"].iloc[-1] == "china_akshare_macro_china_cpi"
    assert len(data["CN_CPI"]) == 24


def test_fetch_us_cpi_with_fred_uses_sa_mom_and_nsa_yoy(monkeypatch) -> None:
    dates = pd.date_range("2025-04-01", periods=13, freq="MS")
    sa_values = [100.0] * 12 + [100.6]
    nsa_values = [100.0] + [101.0] * 11 + [103.8]

    def fake_fetch(series_id: str) -> pd.DataFrame:
        values = sa_values if series_id == cpi_report.US_CPI_SA_SERIES else nsa_values
        return pd.DataFrame({"日期": dates, "指数": values})

    monkeypatch.setattr(cpi_report, "_fetch_fred_series", fake_fetch)

    frame = cpi_report._fetch_us_cpi_with_fred()
    latest = frame.iloc[-1]

    assert round(float(latest["MoM"]), 2) == 0.60
    assert round(float(latest["YoY"]), 2) == 3.80


def test_index_metrics_use_calendar_month_alignment() -> None:
    frame = pd.DataFrame(
        {
            "日期": pd.to_datetime(["2025-03-01", "2025-04-01", "2026-04-01"]),
            "指数": [99.0, 100.0, 103.8],
        }
    )

    metrics = cpi_report._build_metrics_frame_from_index(frame, "test")
    latest = metrics.iloc[-1]

    assert pd.isna(latest["MoM"])
    assert round(float(latest["YoY"]), 2) == 3.80


def test_fetch_hk_cpi_from_api_parses_documented_payload(monkeypatch) -> None:
    class FakeResponse:
        def raise_for_status(self) -> None:
            return None

        def json(self) -> dict:
            return {
                "header": {"status": {"code": 0}},
                "dataSet": [
                    {"freq": "M", "period": "202601", "svDesc": "Index", "figure": 110.0},
                    {"freq": "M", "period": "202601", "svDesc": "Month-to-month % change", "figure": 0.2},
                    {"freq": "M", "period": "202601", "svDesc": "Year-on-year % change", "figure": 1.1},
                ],
            }

    monkeypatch.setattr(cpi_report.requests, "post", lambda *args, **kwargs: FakeResponse())

    frame = cpi_report._fetch_hk_cpi_from_api(period_start="202501")

    assert frame.iloc[-1]["指数"] == 110.0
    assert frame.iloc[-1]["MoM"] == 0.2
    assert frame.iloc[-1]["YoY"] == 1.1
    assert frame.iloc[-1]["数据源"] == "hk_censtatd_api"


def test_plot_cpi_trends_uses_rolling_ten_year_window(monkeypatch, tmp_path: Path) -> None:
    class FakeAxis:
        def __init__(self) -> None:
            self.plots: list[pd.Series] = []

        def plot(self, x, y, label=None, linestyle=None):  # noqa: ANN001
            self.plots.append(pd.Series(pd.to_datetime(x)))

        def set_title(self, *_args, **_kwargs) -> None:
            return None

        def set_xlabel(self, *_args, **_kwargs) -> None:
            return None

        def set_ylabel(self, *_args, **_kwargs) -> None:
            return None

        def legend(self, *_args, **_kwargs) -> None:
            return None

        def grid(self, *_args, **_kwargs) -> None:
            return None

        class Axis:
            def set_major_locator(self, *_args, **_kwargs) -> None:
                return None

            def set_major_formatter(self, *_args, **_kwargs) -> None:
                return None

        xaxis = Axis()

    class FakeFigure:
        def savefig(self, path):  # noqa: ANN001
            Path(path).write_text("ok", encoding="utf-8")

        def tight_layout(self) -> None:
            return None

        def autofmt_xdate(self) -> None:
            return None

    fake_axis = FakeAxis()
    fake_fig = FakeFigure()
    monkeypatch.setattr(cpi_report.plt, "subplots", lambda *args, **kwargs: (fake_fig, fake_axis))
    monkeypatch.setattr(cpi_report.plt, "close", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(cpi_report, "OUTPUT_DIR", tmp_path)

    cpi_report.plot_cpi_trends({"US_CPI": _sample_standard_frame("fred")})

    latest = pd.Timestamp("2026-06-01")
    assert fake_axis.plots
    assert fake_axis.plots[0].min() >= latest - pd.DateOffset(years=10)


def test_crypto_export_adds_generated_timestamp_row(tmp_path: Path) -> None:
    df = pd.DataFrame(
        [
            {
                "Symbol": "BTC",
                "Current_Price": 100.0,
                "Weekly_Volume": 10.0,
                "MoM": 1.0,
                "YoY": 2.0,
                "Pctile_1Y": 30.0,
                "Change_2W": 3.0,
                "Market_Cap": 500.0,
                "Month_Volatility": 0.1,
                "Quarter_Volatility": 0.2,
                "Year_Volatility": 0.3,
                "Month_Sharpe": 1.0,
                "Quarter_Sharpe": 2.0,
                "Year_Sharpe": 3.0,
                "Month_Return": 0.1,
                "Quarter_Return": 0.2,
                "Year_Return": 0.3,
            }
        ]
    )
    target = tmp_path / "crypto_metrics.xlsx"

    assert crypto_report.export_to_excel_template(df, filename=str(target)) is True

    workbook = load_workbook(target)
    sheet = workbook.active
    assert sheet["A1"].value == "指标"
    assert sheet["B1"].value is None
    assert sheet["C1"].value == "BTC"


def test_bond_report_removes_generated_timestamp_row(monkeypatch, tmp_path: Path) -> None:
    monkeypatch.setattr(bond_report, "ROOT", tmp_path)
    monkeypatch.setattr(bond_report, "RAW_DIR", tmp_path / "output" / "raw_data")
    bond_report.RAW_DIR.mkdir(parents=True, exist_ok=True)
    monkeypatch.setattr(bond_report, "_get_china_market_caps", lambda _date: ("1 B CNY", "2 B CNY", pd.DataFrame()))
    monkeypatch.setattr(bond_report, "_get_china_volumes_30d", lambda *_args: ("3 B CNY", "4 B CNY"))
    monkeypatch.setattr(
        bond_report,
        "_get_cn_us_yield_metrics",
        lambda *_args: ((1.0, 2.0, 3.0, 0.4), (1.1, 2.1, 3.1, 0.5), (1.2, 2.2, 3.2, 0.6), (1.3, 2.3, 3.3, 0.7)),
    )
    monkeypatch.setattr(bond_report, "_get_us_market_caps", lambda: ("5 B USD", "6 B USD"))
    monkeypatch.setattr(bond_report, "_resolve_sse_date", lambda *_args: "20260410")

    bond_report.main(debug=False)

    workbook = load_workbook(tmp_path / "output" / "bonds.xlsx")
    sheet = workbook["bonds"]
    assert sheet["A1"].value == "指标类别"
    assert sheet["B1"].value == "中国"


def test_bond_report_uses_cached_us_yields_when_fred_times_out(monkeypatch, tmp_path: Path) -> None:
    monkeypatch.setattr(bond_report, "RAW_DIR", tmp_path)
    dates = pd.date_range("2024-01-02", "2026-04-10", freq="B")
    cn_us = pd.DataFrame(
        {
            "日期": dates,
            "中国国债收益率:2年": [2.0] * len(dates),
            "中国国债收益率:10年": [2.5] * len(dates),
        }
    )
    cached_us = pd.DataFrame({"2Y": [0.04] * len(dates), "10Y": [0.045] * len(dates)}, index=dates)
    cached_us.index.name = "Date"
    cached_us.to_excel(tmp_path / "美债收益率_20260410.xlsx")

    monkeypatch.setattr(bond_report.ak, "bond_zh_us_rate", lambda *args, **kwargs: cn_us)
    monkeypatch.setattr(
        bond_report.requests,
        "get",
        lambda *args, **kwargs: (_ for _ in ()).throw(TimeoutError("fred timeout")),
    )

    _, _, us2, us10 = bond_report._get_cn_us_yield_metrics("20260410")

    assert us2[1] > 0
    assert us10[1] > 0


def test_bond_report_prefers_akshare_us_yields(monkeypatch, tmp_path: Path) -> None:
    monkeypatch.setattr(bond_report, "RAW_DIR", tmp_path)
    dates = pd.date_range("2024-01-02", "2026-04-10", freq="B")
    cn_us = pd.DataFrame(
        {
            "日期": dates,
            "中国国债收益率2年": [2.0] * len(dates),
            "中国国债收益率10年": [2.5] * len(dates),
            "美国国债收益率2年": [4.0] * len(dates),
            "美国国债收益率10年": [4.5] * len(dates),
        }
    )

    monkeypatch.setattr(bond_report.ak, "bond_zh_us_rate", lambda *args, **kwargs: cn_us)
    monkeypatch.setattr(
        bond_report.requests,
        "get",
        lambda *args, **kwargs: (_ for _ in ()).throw(AssertionError("FRED should not be called")),
    )

    _, _, us2, us10 = bond_report._get_cn_us_yield_metrics("20260410")

    assert us2[1] > 0
    assert us10[1] > 0
