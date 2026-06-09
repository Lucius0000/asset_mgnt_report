from __future__ import annotations

from pathlib import Path
import runpy

import pandas as pd
from openpyxl import Workbook, load_workbook

from scripts import overall
from scripts.pipelines import bond_report, crypto_report
from src.asset_mgnt_report.config.defaults import build_app_config
from src.asset_mgnt_report.io.overall_seed_snapshot import (
    get_snapshot_path,
    load_overall_seed_snapshot,
    upsert_sheet1_asset_metrics,
    upsert_sheet2_fx_rows,
)


def _create_overall_seed_template(path: Path) -> None:
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Sheet1"
    ws1.append(
        [
            "区域",
            "资产大类",
            "月收益率年化 (%)",
            "年收益率 (%)",
            "月波动率年化（%）",
            "总市值 ($)",
            "Sharpe Ratio",
            "Adjusted Sharpe Ratio",
            "Gainer",
        ]
    )
    ws1.append(["美国", "股票权益", None, None, None, None, None, None, None])
    ws1.append([None, "债券固收", None, None, None, None, None, None, None])
    ws1.append([None, "房地产", 9.9, 8.8, None, "manual", None, None, "manual"])
    ws1.append(["中国", "股票权益", None, None, None, None, None, None, None])
    ws1.append([None, "债券固收", None, None, None, None, None, None, None])
    ws1.append([None, "房地产", -0.5, -2.8, None, None, None, None, None])
    ws1.append(["香港", "股票权益", None, None, None, None, None, None, None])
    ws1.append([None, "房地产", 0.03, -5.2, None, None, None, None, None])
    ws1.append(["商品与贵金属（黄金）", None, None, None, None, None, None, None, None])
    ws1.append(["数字货币（BTC）", None, None, None, None, None, None, None, None])

    ws2 = wb.create_sheet("Sheet2")
    ws2.append(["货币汇率", "汇率值", "日期", "MoM(%)", "YoY(%)", "5年均值", "5年均值周期"])
    ws2.append(["USD_CNH", 0, "", 0, 0, 0, ""])
    wb.save(path)


def test_snapshot_upsert_overwrites_existing_fields(tmp_path) -> None:
    config = build_app_config(project_root=tmp_path)

    upsert_sheet1_asset_metrics(
        region="美国",
        asset_class="股票权益",
        fields={"月收益率年化 (%)": 0.011, "总市值 ($)": "10 B USD"},
        config=config,
    )
    upsert_sheet1_asset_metrics(
        region="美国",
        asset_class="股票权益",
        fields={"Gainer": "2 B USD", "总市值 ($)": "11 B USD"},
        config=config,
    )

    snapshot = load_overall_seed_snapshot(config)
    record = snapshot["sheet1_assets"]["美国|股票权益"]
    assert record["月收益率年化 (%)"] == 0.011
    assert record["总市值 ($)"] == "11 B USD"
    assert record["Gainer"] == "2 B USD"
    assert get_snapshot_path(config).exists()


def test_overall_main_syncs_seed_and_preserves_manual_rows(tmp_path, monkeypatch) -> None:
    config = build_app_config(project_root=tmp_path)
    input_path = config.seed_data_dir / "整体.xlsx"
    output_path = config.output_dir / "整体_processed.xlsx"
    log_path = config.raw_output_dir / "整体_calculation_steps.txt"
    _create_overall_seed_template(input_path)

    seed_records = {
        ("美国", "股票权益"): {
            "月收益率年化 (%)": 0.4403,
            "年收益率 (%)": 0.2775,
            "月波动率年化（%）": 0.1968,
            "总市值 ($)": "66,538 B USD",
            "Gainer": "4,270.70 B USD",
        },
        ("美国", "债券固收"): {
            "月收益率年化 (%)": 0.038,
            "年收益率 (%)": 0.0364,
            "月波动率年化（%）": 0.0001,
            "总市值 ($)": "38,514 B USD",
            "Gainer": "230 B USD",
        },
        ("中国", "股票权益"): {
            "月收益率年化 (%)": -0.0698,
            "年收益率 (%)": 0.248,
            "月波动率年化（%）": 0.2316,
            "总市值 ($)": "67,791 B CNY",
            "Gainer": "586.42 B CNY",
        },
        ("中国", "债券固收"): {
            "月收益率年化 (%)": 0.0131,
            "年收益率 (%)": 0.0139,
            "月波动率年化（%）": 0.0001,
            "总市值 ($)": "20,304 B CNY",
            "Gainer": "21.42 B CNY",
        },
        ("香港", "股票权益"): {
            "月收益率年化 (%)": -0.101,
            "年收益率 (%)": 0.2808,
            "月波动率年化（%）": 0.2672,
            "总市值 ($)": "49,472 B HKD",
            "Gainer": "1,043.31 B HKD",
        },
        ("商品与贵金属（黄金）", None): {
            "月收益率年化 (%)": -0.6486,
            "年收益率 (%)": 0.3856,
            "月波动率年化（%）": 0.496,
            "总市值 ($)": "33,192.27 B USD",
            "Gainer": "1,874.55 B USD",
        },
        ("数字货币（BTC）", None): {
            "月收益率年化 (%)": 0.310044,
            "年收益率 (%)": -0.119688,
            "月波动率年化（%）": 0.450888,
            "总市值 ($)": "1,490 B USD",
            "Gainer": "128.37 B USD",
        },
    }
    for (region, asset), fields in seed_records.items():
        upsert_sheet1_asset_metrics(region=region, asset_class=asset, fields=fields, config=config)
    upsert_sheet2_fx_rows(
        [
            {"货币汇率": "USD_CNH", "汇率值": 6.8657, "日期": "2026-04-13", "MoM(%)": -0.50719, "YoY(%)": -4.78852, "5年均值": 6.921833, "5年均值周期": "2021-04 to 2026-04"},
            {"货币汇率": "USD_HKD", "汇率值": 7.833444, "日期": "2026-04-13", "MoM(%)": 0.079686, "YoY(%)": 0.973316, "5年均值": 7.809191, "5年均值周期": "2021-04 to 2026-04"},
            {"货币汇率": "CNH_HKD", "汇率值": 1.140953, "日期": "2026-04-13", "MoM(%)": 0.589873, "YoY(%)": 6.051617, "5年均值": 1.130258, "5年均值周期": "2021-04 to 2026-04"},
        ],
        config=config,
    )

    monkeypatch.setattr(overall, "APP_CONFIG", config)
    monkeypatch.setitem(overall.CONFIG, "input_path", input_path)
    monkeypatch.setitem(overall.CONFIG, "output_path", output_path)
    monkeypatch.setitem(overall.CONFIG, "log_path", log_path)

    overall.main()

    synced_seed = load_workbook(input_path)
    synced_sheet1 = synced_seed["Sheet1"]
    assert synced_sheet1["C2"].value == 0.4403
    assert synced_sheet1["C2"].number_format == "0.00%"
    assert synced_sheet1["F2"].value == "66,538 B USD"
    assert synced_sheet1["I2"].value == "4,270.70 B USD"
    assert synced_sheet1["C4"].value == 9.9
    assert synced_sheet1["F4"].value == "manual"
    synced_sheet2 = synced_seed["Sheet2"]
    assert synced_sheet2.max_row == 4
    assert synced_sheet2["A2"].value == "USD_CNH"
    assert synced_sheet2["E4"].value == 6.051617

    processed = load_workbook(output_path)
    processed_sheet1 = processed["Sheet1"]
    assert processed_sheet1["C2"].value == 0.4403
    assert processed_sheet1["C2"].number_format == "0.00%"
    assert processed_sheet1["G2"].value == 2.04
    assert processed_sheet1["H2"].value == 2.04
    assert processed_sheet1["G3"].value == 0
    assert log_path.exists()


def test_overall_main_fails_when_snapshot_is_incomplete(tmp_path, monkeypatch) -> None:
    config = build_app_config(project_root=tmp_path)
    input_path = config.seed_data_dir / "整体.xlsx"
    output_path = config.output_dir / "整体_processed.xlsx"
    log_path = config.raw_output_dir / "整体_calculation_steps.txt"
    _create_overall_seed_template(input_path)

    upsert_sheet1_asset_metrics(
        region="美国",
        asset_class="股票权益",
        fields={"月收益率年化 (%)": 1.0},
        config=config,
    )

    monkeypatch.setattr(overall, "APP_CONFIG", config)
    monkeypatch.setitem(overall.CONFIG, "input_path", input_path)
    monkeypatch.setitem(overall.CONFIG, "output_path", output_path)
    monkeypatch.setitem(overall.CONFIG, "log_path", log_path)

    try:
        overall.main()
    except RuntimeError as exc:
        assert "缺少必要字段" in str(exc)
    else:
        raise AssertionError("overall.main() should fail when snapshot is incomplete")


def test_bond_report_removes_current_time_row(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(bond_report, "ROOT", tmp_path)
    monkeypatch.setattr(bond_report, "RAW_DIR", tmp_path / "output" / "raw_data")
    bond_report.RAW_DIR.mkdir(parents=True, exist_ok=True)
    monkeypatch.setattr(bond_report, "_resolve_sse_date", lambda *_args: "20260411")
    monkeypatch.setattr(bond_report, "_get_china_market_caps", lambda _date: ("1,313 B CNY", "20,304 B CNY", pd.DataFrame()))
    monkeypatch.setattr(bond_report, "_get_china_volumes_30d", lambda *_args: ("72 B CNY", "1,337 B CNY"))
    monkeypatch.setattr(
        bond_report,
        "_get_cn_us_yield_metrics",
        lambda *_args: ((1.31, 1.39, 0.01, 0.0), (1.82, 1.75, 0.01, 0.0), (3.8, 3.64, 0.01, 0.0), (4.31, 4.2, 0.01, 0.0)),
    )
    monkeypatch.setattr(bond_report, "_get_us_market_caps", lambda: ("30,846 B USD", "38,514 B USD"))
    monkeypatch.setattr(bond_report, "upsert_sheet1_asset_metrics", lambda **_: None)

    bond_report.main(debug=False)

    ws = load_workbook(tmp_path / "output" / "bonds.xlsx").active
    assert ws["A1"].value == "指标类别"
    assert ws["B1"].value == "中国"


def test_crypto_report_removes_current_time_row(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(crypto_report, "upsert_sheet1_asset_metrics", lambda **_: None)
    df = pd.DataFrame(
        [
            {
                "Symbol": "BTC",
                "Current_Price": 70000.0,
                "MoM": 0.10,
                "YoY": 0.20,
                "Change_2W": 0.05,
                "Pctile_1Y": 0.80,
                "Market_Cap": 1_500_000_000_000.0,
                "Weekly_Volume": 1_000_000_000.0,
                "Month_Volatility": 0.45,
                "Quarter_Volatility": 0.60,
                "Year_Volatility": 0.43,
                "Month_Sharpe": 0.6,
                "Quarter_Sharpe": 0.5,
                "Year_Sharpe": 0.4,
                "Month_Return": 0.31,
                "Quarter_Return": 0.22,
                "Year_Return": 0.20,
            }
        ]
    )

    output_path = tmp_path / "crypto.xlsx"
    assert crypto_report.export_to_excel_template(df, filename=str(output_path)) is True

    ws = load_workbook(output_path).active
    assert ws["A1"].value == "指标"
    assert ws["C1"].value == "BTC"
    assert ws["C5"].value == 0.10
    assert ws["C5"].number_format == "0.00%"


def test_precious_metals_report_removes_current_time_row(tmp_path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        "src.asset_mgnt_report.config.defaults.build_app_config",
        lambda *args, **kwargs: build_app_config(project_root=tmp_path),
    )
    monkeypatch.setattr("yfinance.Ticker", lambda _symbol: type("FakeTicker", (), {"history": lambda self, period='6y': pd.DataFrame({"Close": [100.0, 101.0], "Volume": [1_000, 1_100]}, index=pd.to_datetime(["2026-04-01", "2026-04-14"]))})())
    monkeypatch.setattr(
        "akshare.spot_golden_benchmark_sge",
        lambda: pd.DataFrame({"交易时间": ["2026-04-01", "2026-04-14"], "晚盘价": [900.0, 920.0]}),
    )
    monkeypatch.setattr(
        "src.asset_mgnt_report.io.overall_seed_snapshot.upsert_sheet1_asset_metrics",
        lambda **_: None,
    )

    runpy.run_module("scripts.pipelines.precious_metals_report", run_name="__main__")

    ws = load_workbook(tmp_path / "output" / "commodity_indicators_summary.xlsx").active
    assert ws["A1"].value == "指标"
